"""Sinh báo cáo Word Phase 5 từ metrics/phase5_metrics.json (mọi số liệu lấy từ file, không gõ tay).

    python -m phase5_ml.build_report --metrics outputs/phase5/<run>/metrics/phase5_metrics.json \
        --charts outputs/phase5/<run>/charts --out docs/BAO_CAO_PHASE5_MLLIB.docx [--reference <metrics.json khác>]
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BLUE = RGBColor(0x1F, 0x4E, 0x8C)


def vi_int(n) -> str:
    return f"{int(n):,}".replace(",", ".")


def vi_num(x, digits=4) -> str:
    if x is None:
        return "–"
    return f"{x:.{digits}f}".replace(".", ",")


def pct(x, digits=1) -> str:
    return f"{x * 100:.{digits}f}".replace(".", ",") + "%"


def shade(cell, color: str):
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:val"), "clear")
    element.set(qn("w:color"), "auto")
    element.set(qn("w:fill"), color)
    props.append(element)


class Report:
    def __init__(self):
        self.doc = Document()
        section = self.doc.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        for side in ("left_margin", "right_margin"):
            setattr(section, side, Cm(2.2))
        section.top_margin = section.bottom_margin = Cm(2)
        normal = self.doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        normal.paragraph_format.space_after = Pt(5)
        for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
            style = self.doc.styles[name]
            style.font.name = "Calibri"
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = BLUE
            style.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        self._footer(section)

    def _footer(self, section):
        paragraph = section.footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("Phase 5 – Spark MLlib · Trang ")
        run.font.size = Pt(9)
        for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
            r = paragraph.add_run()
            r.font.size = Pt(9)
            if tag:
                fld = OxmlElement("w:fldChar")
                fld.set(qn("w:fldCharType"), tag)
                r._r.append(fld)
            else:
                instr = OxmlElement("w:instrText")
                instr.set(qn("xml:space"), "preserve")
                instr.text = text
                r._r.append(instr)

    def h1(self, text):
        self.doc.add_heading(text, level=1)

    def h2(self, text):
        self.doc.add_heading(text, level=2)

    def p(self, text="", bold=False, italic=False, size=None, align=None, color=None):
        paragraph = self.doc.add_paragraph()
        run = paragraph.add_run(text)
        run.bold, run.italic = bold, italic
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
        if align:
            paragraph.alignment = align
        return paragraph

    def rich(self, parts):
        """parts: list[(text, bold)]"""
        paragraph = self.doc.add_paragraph()
        for text, bold in parts:
            paragraph.add_run(text).bold = bold
        return paragraph

    def bullets(self, items):
        for item in items:
            paragraph = self.doc.add_paragraph(style="List Bullet")
            if isinstance(item, tuple):
                paragraph.add_run(item[0]).bold = True
                paragraph.add_run(item[1])
            else:
                paragraph.add_run(item)
            paragraph.paragraph_format.space_after = Pt(2)

    def table(self, header, rows, widths_cm, font=9.5, bold_rows=()):
        table = self.doc.add_table(rows=1, cols=len(header))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        table._tbl.tblPr.append(layout)
        for i, text in enumerate(header):
            cell = table.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(text)
            run.bold = True
            run.font.size = Pt(font)
            shade(cell, "D9E2F3")
        for r_index, row in enumerate(rows):
            cells = table.add_row().cells
            for i, text in enumerate(row):
                cells[i].text = ""
                run = cells[i].paragraphs[0].add_run(str(text))
                run.font.size = Pt(font)
                run.bold = r_index in bold_rows
                if i > 0 and re.fullmatch(r"[\d.,%–\- ]+", str(text)):
                    cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for i, width in enumerate(widths_cm):
            table.columns[i].width = Cm(width)
        for row in table.rows:
            for i, width in enumerate(widths_cm):
                row.cells[i].width = Cm(width)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def image(self, path, width_cm=16.0, caption=None):
        self.doc.add_picture(str(path), width=Cm(width_cm))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if caption:
            self.p(caption, italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    def save(self, path):
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(path)


def build(metrics_path: str, charts_dir: str, out_path: str, reference_path: str | None = None):
    m = json.loads(Path(metrics_path).read_text(encoding="utf-8"))
    ref = json.loads(Path(reference_path).read_text(encoding="utf-8")) if reference_path else None
    charts = Path(charts_dir)
    k = m["config"]["top_k"]
    data, split, cv, held = m["data"], m["split"], m["cross_validation"], m["heldout_test"]
    rmse, serving, timing = m["rmse"], m["serving_model"], m["timings_seconds"]
    sel = cv["selected_parameters"]
    quick = m["config"]["quick_grid"]
    n_candidates = len(cv["candidates"])
    als_w, pop_w, hyb_w = held["als"]["warm"], held["popularity"]["warm"], held["hybrid_als_popularity"]["warm"]
    als_o, pop_o, hyb_o = (held[x]["overall_weighted"] for x in ("als", "popularity", "hybrid_als_popularity"))
    cold_share = held["cold_user_share"]
    rt = rmse["heldout_test"]["rmse"]

    r = Report()
    # ---------------- Trang bìa -----------------------------------------------------------------
    r.p("BÀI TẬP NHÓM – HỆ THỐNG BIG DATA", bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, color=BLUE)
    r.p("BÁO CÁO PHASE 5", bold=True, size=22, align=WD_ALIGN_PARAGRAPH.CENTER, color=BLUE)
    r.p("Tầng phân tích / ML – Spark MLlib", bold=True, size=15, align=WD_ALIGN_PARAGRAPH.CENTER)
    r.p("Hệ thống phân tích hành vi người dùng thương mại điện tử & gợi ý sản phẩm", italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER)
    r.table(["Mục", "Nội dung"], [
        ["Thành viên phụ trách", "Phan Thị Mỹ Duyên – 23139008"],
        ["Phần việc (ActionPlan)", "Phase 5 – task 16, 17, 18"],
        ["Dữ liệu", "RetailRocket (Kaggle) – " + vi_int(data["raw_rows"]) + " sự kiện, "
                    + data["first_event_utc"][:10] + " → " + data["last_event_utc"][:10]],
        ["Mã chạy", m["run_id"] + " · Spark " + m["environment"]["spark"]],
    ], [4.6, 11.8], font=10)

    # ---------------- 1. Tóm tắt -----------------------------------------------------------------
    r.h1("1. Tóm tắt kết quả")
    r.table(["Task", "Nội dung", "Kết quả", "Trạng thái"], [
        ["16", "Pipeline MLlib (VectorAssembler + ALS)", "Pipeline 2 stage, đã lưu", "Hoàn thành"],
        ["17", "Huấn luyện offline trên Delta Lake (Silver/Gold)",
         f"{vi_int(data['clean_rows'])} sự kiện, {vi_int(serving['interaction_pairs'])} cặp user–item", "Hoàn thành*"],
        ["18", "Đánh giá: RMSE / Precision@K / cross-validation",
         f"CV 2 fold thời gian × {n_candidates} cấu hình; held-out; RMSE", "Hoàn thành"],
    ], [1.4, 6.2, 6.6, 2.2])
    r.p("* Mã đọc được Silver Delta của nhóm (--source delta). Lần chạy này dùng Parquet cùng schema Silver/Gold "
        "vì môi trường chạy không tải được JAR Delta (Maven bị chặn); chưa chạy trực tiếp trên HDFS của nhóm.",
        italic=True, size=9)
    uplift_ndcg = als_w["ndcg_at_k"] / pop_w["ndcg_at_k"] - 1
    uplift_p = als_w["precision_at_k"] / pop_w["precision_at_k"]
    r.bullets([
        ("Mô hình chọn: ", f"ALS implicit, rank {sel['rank']}, regParam {sel['reg_param']}, alpha {sel['alpha']:g}."),
        (f"User đã có lịch sử (warm), held-out: ",
         f"ALS đạt NDCG@{k} = {vi_num(als_w['ndcg_at_k'])} so với Popularity {vi_num(pop_w['ndcg_at_k'])} "
         f"(+{pct(uplift_ndcg, 0)}); Precision@{k} gấp {vi_num(uplift_p, 1)} lần."),
        ("Vấn đề lớn nhất là cold-start: ", f"{pct(cold_share)} user ở giai đoạn test chưa có lịch sử trong train. "
         "ALS không gợi ý được cho nhóm này."),
        ("Toàn bộ user (gồm cold): ", f"Hybrid (ALS cho warm, Popularity cho cold) NDCG@{k} = {vi_num(hyb_o['ndcg_at_k'])}, "
         f"Popularity = {vi_num(pop_o['ndcg_at_k'])}, ALS riêng = {vi_num(als_o['ndcg_at_k'])}."),
        ("RMSE (ALS explicit, mức hành vi 1/3/5): ", f"{vi_num(rt['als_clipped_1_5'], 3)} so với baseline trung bình toàn cục "
         f"{vi_num(rt['baseline_global_mean'], 3)} – ALS không vượt baseline (xem mục 4.4)."),
    ])

    # ---------------- 2. Task 16 -----------------------------------------------------------------
    r.h1("2. Task 16 – Pipeline MLlib")
    r.image(charts / "00_pipeline.png", 16.4, "Hình 1. Luồng Phase 5 trong pipeline tổng thể")
    r.p("Pipeline Spark ML gồm hai stage, dùng chung một đối tượng Pipeline (fit/save/load như một khối):")
    r.table(["Stage", "Thành phần", "Cấu hình chính"], [
        ["0", "VectorAssembler", "view_count, cart_count, purchase_count → interaction_features (vector đặc trưng phụ, lưu cùng pipeline)"],
        ["1", "ALS (implicit)", f"userCol = visitorid, itemCol = itemid, ratingCol = rating; implicitPrefs = True; rank {sel['rank']}; "
                                f"regParam {sel['reg_param']}; alpha {sel['alpha']:g}; seed 42; coldStartStrategy = drop"],
    ], [1.5, 3.6, 11.3])
    r.bullets([
        ("Tín hiệu: ", "rating = view + 3 × addtocart + 5 × transaction (trùng điểm Gold của Phase 4: 1/3/5). "
                       "Đã kiểm tra từng cặp: rating ALS = interaction_score của Gold, " + vi_int(m['gold']['mismatched_pairs']) + " cặp lệch."),
        ("Vì sao ALS, không phải Random Forest: ", "dữ liệu là phản hồi ngầm (chỉ có hành vi, không có điểm đánh giá), ma trận user–item rất thưa; "
         "ALS học nhân tố ẩn trực tiếp cho bài toán gợi ý. Random Forest cần nhãn và đặc trưng cố định, không sinh được danh sách top-K cho từng user."),
        ("Lưu ý trung thực: ", "ALS nhận ba cột (user, item, rating), không dùng vector interaction_features; vector này là đặc trưng phụ để kiểm tra/mở rộng."),
    ])

    # ---------------- 3. Task 17 -----------------------------------------------------------------
    r.h1("3. Task 17 – Huấn luyện offline")
    r.h2("3.1 Dữ liệu đầu vào")
    ev = data["event_counts"]
    r.table(["Chỉ tiêu", "Giá trị"], [
        ["Dòng gốc / sau làm sạch", f"{vi_int(data['raw_rows'])} / {vi_int(data['clean_rows'])} (bỏ {vi_int(data['removed_rows'])} dòng trùng)"],
        ["view / addtocart / transaction", f"{vi_int(ev['view'])} / {vi_int(ev['addtocart'])} / {vi_int(ev['transaction'])}"],
        ["Số user / số item", f"{vi_int(data['users'])} / {vi_int(data['items'])}"],
        ["Cặp user–item (Gold)", vi_int(m["gold"]["gold_pairs"])],
        ["Khoảng thời gian", f"{data['first_event_utc']} → {data['last_event_utc']}"],
    ], [5.6, 10.8])
    r.p("Làm sạch theo đúng quy tắc Silver của Phase 4: bỏ dòng thiếu visitorid/itemid, chỉ giữ 3 loại sự kiện, loại trùng theo 5 cột gốc.")
    r.h2("3.2 Chia dữ liệu theo thời gian")
    r.image(charts / "01_split_timeline.png", 13.5, "Hình 2. Số sự kiện từng phần (mốc 50% / 70% / 85% theo timestamp)")
    ep = split["events_per_partition"]
    r.p("Chia theo timestamp của sự kiện thô trước khi gộp user–item, nên không có thông tin tương lai rò vào tập train. "
        f"Mốc: {split['boundaries_utc'][0][:10]}, {split['boundaries_utc'][1][:10]}, {split['boundaries_utc'][2][:10]}.")
    r.table(["Bước", "Dữ liệu fit", "Dữ liệu đánh giá"], [
        ["Fold 1", "đến mốc 50%", "(50%, 70%]"],
        ["Fold 2", "đến mốc 70%", "(70%, 85%]"],
        ["Held-out test", f"đến mốc 85% ({vi_int(held['train_events'])} sự kiện)", f"> 85% ({vi_int(held['test_events'])} sự kiện)"],
        ["Model phục vụ", "toàn bộ dữ liệu (fit sau khi đã chốt metric)", "–"],
    ], [3.4, 7.0, 6.0])
    r.h2("3.3 Quy trình huấn luyện & sản phẩm")
    r.bullets([
        f"Chế độ {'nhanh' if quick else 'đầy đủ'}: {n_candidates} cấu hình ALS × 2 fold, {cv['iterations']} vòng lặp khi chọn tham số, "
        f"{held['fit_iterations']} vòng lặp khi fit model đánh giá/phục vụ.",
        f"Thời gian (máy 2 nhân, Spark local): nạp + Gold {timing['1_load_clean_gold']:.0f} s · cross-validation {timing['3_cross_validation']:.0f} s · "
        f"held-out {timing['4_heldout_test']:.0f} s · RMSE {timing['4b_rmse']:.0f} s · model phục vụ {timing['5_serving_model']:.0f} s · "
        f"tổng {timing['total'] / 60:.1f} phút.",
        f"Model phục vụ: {vi_int(serving['trained_users'])} user, {vi_int(serving['catalog_size'])} item, fit {serving['fit_seconds']:.0f} s.",
        "Sản phẩm: models/als_pipeline (PipelineModel), models/evaluation_pipeline, gợi ý top-10 mẫu "
        f"({vi_int(serving['sample_users_exported'])} user) dạng Parquet + CSV, popular_items_top50.csv, lake/gold_interactions.",
    ])

    # ---------------- 4. Task 18 -----------------------------------------------------------------
    r.h1("4. Task 18 – Đánh giá mô hình")
    r.h2("4.1 Cách đo")
    r.bullets([
        f"Mục tiêu: cặp (user, item) MỚI trong cửa sổ held-out. Item user đã xem trong train bị loại khỏi gợi ý (đúng hành vi API).",
        f"Precision@{k} = số item đúng / {k}; Recall@{k} = số item đúng / số item mới của user; NDCG@{k} nhị phân. Trung bình macro theo user.",
        "Warm = user có lịch sử trong train; cold = user mới. Cold user không bị loại: ALS không gợi ý được thì tính là trượt (0).",
        "Lấy mẫu phân tầng (hash cố định, seed 42) để chạy trên máy 2 nhân; số liệu tổng hợp gia quyền lại theo quy mô quần thể thật của từng tầng.",
        "Baseline: Popularity (top item theo tổng rating trong train, loại item đã xem). Hybrid: ALS cho warm, Popularity cho cold.",
    ])
    r.h2("4.2 Cross-validation theo thời gian")
    rows = []
    for c in cv["candidates"]:
        p = c["parameters"]
        mark = "  ← chọn" if p == sel else ""
        rows.append([f"{p['rank']} / {p['reg_param']} / {p['alpha']:g}{mark}",
                     vi_num(c["folds"][0]["als"]["warm"]["ndcg_at_k"]), vi_num(c["folds"][1]["als"]["warm"]["ndcg_at_k"]),
                     vi_num(c["mean_warm_ndcg_at_k"]), vi_num(c["mean_warm_recall_at_k"]), vi_num(c["mean_warm_precision_at_k"])])
    best_index = next(i for i, c in enumerate(cv["candidates"]) if c["parameters"] == sel)
    r.table(["rank / reg / alpha", f"NDCG@{k} fold 1", f"NDCG@{k} fold 2", f"NDCG@{k} TB", f"Recall@{k} TB", f"Precision@{k} TB"],
            rows, [4.6, 2.3, 2.3, 2.5, 2.4, 2.3], bold_rows=(best_index,))
    r.p("Tiêu chí chọn: NDCG@K trung bình của user warm trên hai fold (cold không có tín hiệu để phân biệt cấu hình).", italic=True, size=9.5)
    r.image(charts / "02_cross_validation.png", 14.5, "Hình 3. NDCG@K (user warm) theo cấu hình và fold")
    max_rank = max(c["parameters"]["rank"] for c in cv["candidates"])
    min_rank = min(c["parameters"]["rank"] for c in cv["candidates"])
    if max_rank > min_rank and sel["rank"] == max_rank:
        by_rank = {}
        for c in cv["candidates"]:
            by_rank.setdefault(c["parameters"]["rank"], []).append(c["mean_warm_ndcg_at_k"])
        lo, hi = max(by_rank[min_rank]), max(by_rank[max_rank])
        r.p(f"Rank là yếu tố ảnh hưởng nhất: rank {max_rank} đạt NDCG@{k} tốt nhất {vi_num(hi)} so với {vi_num(lo)} của rank {min_rank} "
            f"(+{pct(hi / lo - 1, 0)}); reg và alpha chỉ thay đổi nhỏ. Cấu hình chọn nằm ở biên của lưới (rank {max_rank}) nên chưa chắc là tối ưu – "
            f"nên thử rank cao hơn (ví dụ 32, 64).", bold=False)
    fold_pop = [f"{vi_num(c['popularity']['warm']['ndcg_at_k'])}" for c in cv["candidates"][0]["folds"]]
    r.p(f"Popularity trên cùng hai fold: NDCG@{k} = {fold_pop[0]} và {fold_pop[1]} (user warm). ALS cao hơn ở cả hai fold.")

    r.h2("4.3 Held-out test (train ≤ 85%, test > 85%)")
    r.p(f"{vi_int(held['eligible_users'])} user có item mới trong giai đoạn test: {vi_int(held['eligible_warm_users'])} warm "
        f"và {vi_int(held['eligible_cold_users'])} cold ({pct(cold_share)}).")
    rows = []
    for key, name in (("als", "ALS"), ("popularity", "Popularity"), ("hybrid_als_popularity", "Hybrid")):
        w, o = held[key]["warm"], held[key]["overall_weighted"]
        rows.append([name, vi_num(w["precision_at_k"]), vi_num(w["recall_at_k"]), vi_num(w["ndcg_at_k"]), pct(w["hit_rate_at_k"], 2),
                     vi_num(o["precision_at_k"]), vi_num(o["recall_at_k"]), vi_num(o["ndcg_at_k"])])
    r.table(["Mô hình", f"P@{k} warm", f"R@{k} warm", f"NDCG warm", "Hit warm", f"P@{k} tổng", f"R@{k} tổng", "NDCG tổng"],
            rows, [2.6, 2.0, 2.0, 2.0, 1.9, 2.0, 2.0, 2.0], font=9)
    r.image(charts / "03_heldout_ranking.png", 16.4, "Hình 4. Held-out: hàng trên – user warm; hàng dưới – toàn bộ user mục tiêu (gồm cold)")
    cov = held["als"].get("catalog_coverage_warm_sample")
    r.p(f"ALS gợi ý {pct(cov, 2)} catalog train cho nhóm warm (khoảng {vi_int(round(cov * held['train_catalog_size']))} item khác nhau): "
        "mô hình còn thiên về item phổ biến.")

    r.h2("4.4 RMSE (phụ trợ, theo yêu cầu ActionPlan)")
    r.p("RMSE chỉ có nghĩa với nhãn số. Mô hình chính là implicit nên không có RMSE tự nhiên; ở đây huấn luyện thêm ALS explicit "
        "với nhãn = mức hành vi mạnh nhất của cặp (view = 1, addtocart = 3, transaction = 5), chỉ chấm cặp mới mà user/item đã biết, "
        "kẹp dự đoán về [1, 5], so với hai baseline trên cùng tập cặp.")
    sel_grid = next(g for g in rmse["grid"] if g["reg_param"] == rmse["selected"]["reg_param"])
    rows = []
    for name, res in (("Fold 1", sel_grid["folds"][0]), ("Fold 2", sel_grid["folds"][1]), ("Test", rmse["heldout_test"])):
        v = res["rmse"]
        rows.append([name, pct(res["coverage"]), vi_num(v["als_clipped_1_5"], 3), vi_num(v["baseline_item_mean"], 3),
                     vi_num(v["baseline_global_mean"], 3)])
    r.table(["Cửa sổ", "Phủ cặp mới", "ALS explicit", "TB theo item", "TB toàn cục"], rows, [3.2, 3.0, 3.4, 3.4, 3.4])
    r.image(charts / "04_rmse.png", 13.5, f"Hình 5. RMSE (regParam {rmse['selected']['reg_param']}, rank {rmse['selected']['rank']})")
    r.p(f"ALS explicit không vượt baseline trung bình (test: {vi_num(rt['als_clipped_1_5'], 3)} so với {vi_num(rt['baseline_global_mean'], 3)}). "
        f"Lý do: nhãn rất lệch (trung bình {vi_num(rmse['heldout_test']['test_mean_level'], 2)}, phần lớn là view = 1) và chỉ "
        f"{pct(rmse['heldout_test']['coverage'])} cặp mới có cả user lẫn item đã biết. Vì vậy RMSE không dùng để chọn mô hình; "
        f"chất lượng gợi ý được đánh giá bằng chỉ số xếp hạng ở mục 4.3.")

    r.h2("4.5 Nhận xét")
    hybrid_gain = hyb_o["ndcg_at_k"] / pop_o["ndcg_at_k"] - 1
    r.bullets([
        f"ALS có giá trị thật với user có lịch sử: NDCG@{k} cao hơn Popularity {pct(uplift_ndcg, 0)}, ổn định ở cả hai fold và held-out.",
        f"Giá trị tuyệt đối thấp (NDCG@{k} ≈ {vi_num(als_w['ndcg_at_k'], 3)}) là đặc trưng của RetailRocket: gần 97% là lượt view, "
        "mỗi user chỉ vài tương tác, đoạn test chủ yếu là item/user chưa từng xuất hiện.",
        f"Hybrid chỉ hơn Popularity {pct(hybrid_gain)} trên toàn bộ user vì {pct(cold_share)} user là cold – "
        "đúng thiết kế API của nhóm (ALS + popularity fallback); cần retrain thường xuyên (Airflow) để user mới sớm trở thành warm.",
        "Hướng cải thiện: thêm đặc trưng nội dung item (category/thuộc tính, dataset có item_properties), thử rank cao hơn lưới hiện tại, "
        "đánh giá trên cửa sổ dài hơn, đưa session/window feature từ Phase 3 cho cold-start.",
    ])

    # ---------------- 5. Hạn chế ------------------------------------------------------------------
    r.h1("5. Hạn chế & lưu ý trung thực")
    r.bullets([
        "Delta Lake: lần chạy này đọc Parquet (cùng schema Silver/Gold) do môi trường không tải được JAR Delta; mã hỗ trợ --source delta để chạy trong Docker của nhóm.",
        f"Đánh giá dùng mẫu user có seed cố định: warm {vi_int(held['als']['warm']['users_evaluated'])} user (gần như toàn bộ), "
        f"cold {vi_int(held['als']['cold']['users_evaluated'])} / {vi_int(held['eligible_cold_users'])} user; số tổng hợp đã gia quyền lại.",
        f"Gợi ý top-{k} xuất cho mẫu {vi_int(serving['sample_users_exported'])} user; xuất toàn bộ {vi_int(serving['trained_users'])} user cần cụm lớn hơn "
        "hoặc pipeline/train.py của nhóm (đã có bucket hóa ứng viên).",
        "Chỉ số offline không chứng minh tăng doanh thu/CTR; dataset không có giá, impression hay review thật.",
        "Mô hình không tự học từng sự kiện; cập nhật bằng huấn luyện lại định kỳ.",
    ])
    if ref is not None:
        rw = ref["heldout_test"]
        r.p(f"Đối chiếu: lần chạy nhanh ({ref['run_id']}, {len(ref['cross_validation']['candidates'])} cấu hình, "
            f"rank {ref['cross_validation']['selected_parameters']['rank']}) có NDCG@{k} warm held-out = "
            f"{vi_num(rw['als']['warm']['ndcg_at_k'])}.", italic=True, size=9.5)

    # ---------------- 6. Chạy lại -----------------------------------------------------------------
    r.h1("6. Chạy lại & tệp bàn giao")
    r.table(["Việc", "Lệnh"], [
        ["Kiểm thử (Spark local)", "python -m unittest discover -s tests -p test_phase5_ml.py -v"],
        ["Chạy toàn bộ (CSV)", "python -m phase5_ml.run_phase5 --source csv --input data/raw/events.csv --run-id p5-01 [--full-grid]"],
        ["Chạy trên Delta Silver", "python -m phase5_ml.run_phase5 --source delta --input hdfs://namenode:8020/ecommerce/delta/silver/events --run-id p5-delta-01"],
        ["Vẽ biểu đồ", "python -m phase5_ml.charts outputs/phase5/<run>/metrics/phase5_metrics.json outputs/phase5/<run>/charts"],
    ], [4.2, 12.2], font=9)
    r.bullets([
        ("phase5_ml/ ", "data.py (nạp/làm sạch/Gold), evaluate.py (Top-K, baseline, hybrid), explicit.py (RMSE), run_phase5.py, charts.py, build_report.py."),
        ("pipeline/ ", "train.py, ranking.py, config.py – lấy nguyên từ repo của nhóm (ALS pipeline, chia thời gian)."),
        ("phase5_results/ ", "metrics/phase5_metrics.json, charts/*.png, recommendations_top10_sample.csv, popular_items_top50.csv (model và Gold Parquet nặng nên không đưa lên GitHub; chạy lại để tạo)."),
    ])
    r.save(out_path)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--metrics", required=True)
    parser.add_argument("--charts", required=True)
    parser.add_argument("--out", required=True)
    parser.add_argument("--reference")
    args = parser.parse_args()
    build(args.metrics, args.charts, args.out, args.reference)
    print("saved", args.out)
